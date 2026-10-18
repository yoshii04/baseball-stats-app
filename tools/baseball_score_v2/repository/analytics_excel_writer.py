"""Analytics Excel writer: writes analysis sheets to Excel file"""
import logging
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.utils import get_column_letter


class AnalyticsExcelWriter:
    # Style constants
    _HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
    _HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
    _ALT_ROW_FILL = PatternFill("solid", fgColor="DCE6F1")
    _BLOCK_HEADER_FILL = PatternFill("solid", fgColor="2E74B5")
    _BLOCK_HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
    _THIN_BORDER = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    def __init__(self, output_path: str):
        self.output_path = output_path
        self.logger = logging.getLogger("app.repository")
        self.wb = Workbook()
        # Remove default empty sheet
        if "Sheet" in self.wb.sheetnames:
            del self.wb["Sheet"]

    # ------------------------------------------------------------------ helpers

    def _apply_header_style(self, ws, row: int, col_count: int) -> None:
        for col in range(1, col_count + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = self._HEADER_FONT
            cell.fill = self._HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 24

    def _apply_borders(self, ws, min_row: int, max_row: int, min_col: int, max_col: int) -> None:
        for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
            for cell in row:
                cell.border = self._THIN_BORDER

    # ------------------------------------------------------------------ Sheet 1

    def write_player_batting_stats(self, stats: list[dict]) -> None:
        """Sheet 1: 選手別打撃成績"""
        ws = self.wb.create_sheet("選手別打撃成績")

        headers = [
            "選手名", "試合数", "打席", "打数", "安打", "単打",
            "2塁打", "3塁打", "本塁打", "打点", "四死球", "三振", "盗塁",
            "打率", "出塁率", "長打率", "OPS",
        ]
        NUM_COLS = len(headers)
        RATE_COLS = [14, 15, 16, 17]  # 打率, 出塁率, 長打率, OPS (1-indexed)
        OPS_COL = 17

        # Header row
        for col, h in enumerate(headers, 1):
            ws.cell(row=1, column=col, value=h)
        self._apply_header_style(ws, 1, NUM_COLS)

        # Data rows
        for i, stat in enumerate(stats):
            row = i + 2
            ws.cell(row=row, column=1, value=stat["player_name"])
            ws.cell(row=row, column=2, value=int(stat["game_count"]))
            ws.cell(row=row, column=3, value=int(stat["at_bat"]))
            ws.cell(row=row, column=4, value=int(stat["at_bat_count"]))
            ws.cell(row=row, column=5, value=int(stat["hit_count"]))
            ws.cell(row=row, column=6, value=int(stat["single"]))
            ws.cell(row=row, column=7, value=int(stat["double_hit"]))
            ws.cell(row=row, column=8, value=int(stat["triple"]))
            ws.cell(row=row, column=9, value=int(stat["homerun"]))
            ws.cell(row=row, column=10, value=int(stat["rbi"]))
            ws.cell(row=row, column=11, value=int(stat["walk_count"]))
            ws.cell(row=row, column=12, value=int(stat["strike_out"]))
            ws.cell(row=row, column=13, value=int(stat["stolen_base"]))
            ws.cell(row=row, column=14, value=float(stat["batting_avg"]) if stat["batting_avg"] is not None else None)
            ws.cell(row=row, column=15, value=float(stat["on_base_pct"]) if stat["on_base_pct"] is not None else None)
            ws.cell(row=row, column=16, value=float(stat["slg"]) if stat["slg"] is not None else None)
            ws.cell(row=row, column=17, value=float(stat["ops"]) if stat["ops"] is not None else None)

            for col in RATE_COLS:
                ws.cell(row=row, column=col).number_format = "0.000"

            # Alternating row color
            if i % 2 == 1:
                for col in range(1, NUM_COLS + 1):
                    ws.cell(row=row, column=col).fill = self._ALT_ROW_FILL

        last_data_row = 1 + len(stats)

        # Borders
        self._apply_borders(ws, 1, last_data_row, 1, NUM_COLS)

        # OPS color scale (red=low → yellow=mid → green=high)
        if len(stats) > 1:
            ops_letter = get_column_letter(OPS_COL)
            ws.conditional_formatting.add(
                f"{ops_letter}2:{ops_letter}{last_data_row}",
                ColorScaleRule(
                    start_type="min", start_color="FF4444",
                    mid_type="percentile", mid_value=50, mid_color="FFFF00",
                    end_type="max", end_color="44BB44",
                ),
            )

        # Freeze header
        ws.freeze_panes = "A2"

        # Column widths
        ws.column_dimensions["A"].width = 12
        for col in range(2, 14):
            ws.column_dimensions[get_column_letter(col)].width = 7
        for col in range(14, 18):
            ws.column_dimensions[get_column_letter(col)].width = 8

        # Chart 1: OPS ランキング（横棒グラフ）
        if stats:
            chart1 = BarChart()
            chart1.type = "bar"
            chart1.title = "OPSランキング"
            chart1.y_axis.title = "選手名"
            chart1.x_axis.title = "OPS"
            chart1.width = 18
            chart1.height = max(8, len(stats) * 0.9)

            data1 = Reference(ws, min_col=OPS_COL, min_row=1, max_row=last_data_row)
            cats1 = Reference(ws, min_col=1, min_row=2, max_row=last_data_row)
            chart1.add_data(data1, titles_from_data=True)
            chart1.set_categories(cats1)
            ws.add_chart(chart1, "S2")

            # Chart 2: 安打種別積み上げ棒グラフ（単打/2塁打/3塁打/本塁打）
            chart2 = BarChart()
            chart2.type = "col"
            chart2.grouping = "stacked"
            chart2.overlap = 100
            chart2.title = "安打種別内訳"
            chart2.x_axis.title = "選手名"
            chart2.y_axis.title = "本数"
            chart2.width = 18
            chart2.height = 12

            data2 = Reference(ws, min_col=6, max_col=9, min_row=1, max_row=last_data_row)
            cats2 = Reference(ws, min_col=1, min_row=2, max_row=last_data_row)
            chart2.add_data(data2, titles_from_data=True)
            chart2.set_categories(cats2)
            ws.add_chart(chart2, "S22")

        self.logger.info("Sheet 1 書き込み完了: %d 選手", len(stats))

    # ------------------------------------------------------------------ Sheet 2

    def write_hit_direction(self, rows: list[dict]) -> None:
        """Sheet 2: 打球方向分析"""
        ws = self.wb.create_sheet("打球方向分析")

        headers = [
            "選手名", "右(右9)", "右中", "中", "左中", "左(左7)", "合計",
            "%右9", "%右中", "%中", "%左中", "%左7",
        ]
        NUM_COLS = len(headers)
        DIR_COLS = ["B", "C", "D", "E", "F"]

        for col, h in enumerate(headers, 1):
            ws.cell(row=1, column=col, value=h)
        self._apply_header_style(ws, 1, NUM_COLS)

        for i, stat in enumerate(rows):
            row = i + 2
            ws.cell(row=row, column=1, value=stat["player_name"])
            ws.cell(row=row, column=2, value=int(stat["right_9"]))
            ws.cell(row=row, column=3, value=int(stat["right_center"]))
            ws.cell(row=row, column=4, value=int(stat["center_field"]))
            ws.cell(row=row, column=5, value=int(stat["left_center"]))
            ws.cell(row=row, column=6, value=int(stat["left_7"]))
            ws.cell(row=row, column=7, value=int(stat["total"]))
            ws.cell(row=row, column=8, value=stat["pct_right_9"])
            ws.cell(row=row, column=9, value=stat["pct_right_center"])
            ws.cell(row=row, column=10, value=stat["pct_center"])
            ws.cell(row=row, column=11, value=stat["pct_left_center"])
            ws.cell(row=row, column=12, value=stat["pct_left_7"])

            for col in range(8, 13):
                ws.cell(row=row, column=col).number_format = "0.0"

            if i % 2 == 1:
                for col in range(1, NUM_COLS + 1):
                    ws.cell(row=row, column=col).fill = self._ALT_ROW_FILL

        last_data_row = 1 + len(rows)
        self._apply_borders(ws, 1, last_data_row, 1, NUM_COLS)

        # Data bars on direction count columns
        for col_letter in DIR_COLS:
            ws.conditional_formatting.add(
                f"{col_letter}2:{col_letter}{last_data_row}",
                DataBarRule(
                    start_type="min", start_value=0,
                    end_type="max", end_value=None,
                    color="4472C4",
                ),
            )

        ws.freeze_panes = "A2"
        ws.column_dimensions["A"].width = 12
        for col in range(2, NUM_COLS + 1):
            ws.column_dimensions[get_column_letter(col)].width = 8

        # Chart: 100%積み上げ棒グラフ
        if rows:
            chart = BarChart()
            chart.type = "col"
            chart.grouping = "percentStacked"
            chart.overlap = 100
            chart.title = "打球方向分布（100%積み上げ）"
            chart.x_axis.title = "選手名"
            chart.width = 20
            chart.height = 14

            data = Reference(ws, min_col=2, max_col=6, min_row=1, max_row=last_data_row)
            cats = Reference(ws, min_col=1, min_row=2, max_row=last_data_row)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            ws.add_chart(chart, "N2")

        self.logger.info("Sheet 2 書き込み完了: %d 選手", len(rows))

    # ------------------------------------------------------------------ Sheet 3

    def write_player_progress(self, rows: list[dict]) -> None:
        """Sheet 3: 選手別成績推移（選手ごとのブロック + 累積打率折れ線グラフ）"""
        ws = self.wb.create_sheet("選手別成績推移")

        col_headers = [
            "日付", "対戦相手", "試合", "打数", "安打",
            "2塁打", "3塁打", "本塁打", "四死球", "三振", "盗塁",
            "累積打率", "累積OPS",
        ]
        NUM_COLS = len(col_headers)

        # Group by player (order preserved from service — sorted by player_name)
        player_rows: dict[str, list[dict]] = defaultdict(list)
        for row in rows:
            player_rows[row["player_name"]].append(row)

        current_row = 1
        chart_series_info: list[dict] = []  # for line chart

        for player_name in sorted(player_rows.keys()):
            games = player_rows[player_name]

            # Player block header
            ws.merge_cells(
                start_row=current_row, start_column=1,
                end_row=current_row, end_column=NUM_COLS,
            )
            cell = ws.cell(row=current_row, column=1, value=f"【{player_name}】")
            cell.font = self._BLOCK_HEADER_FONT
            cell.fill = self._BLOCK_HEADER_FILL
            cell.alignment = Alignment(horizontal="left", vertical="center")
            ws.row_dimensions[current_row].height = 20
            current_row += 1

            # Column headers
            for col, h in enumerate(col_headers, 1):
                ws.cell(row=current_row, column=col, value=h)
            self._apply_header_style(ws, current_row, NUM_COLS)
            current_row += 1

            # Data rows
            data_start = current_row
            for game in games:
                ws.cell(row=current_row, column=1, value=str(game["game_date"]))
                ws.cell(row=current_row, column=2, value=game["opponent"])
                ws.cell(row=current_row, column=3, value=int(game["game_number"]))
                ws.cell(row=current_row, column=4, value=int(game["at_bat_count"]))
                ws.cell(row=current_row, column=5, value=int(game["hit_count"]))
                ws.cell(row=current_row, column=6, value=int(game["double_hit"]))
                ws.cell(row=current_row, column=7, value=int(game["triple"]))
                ws.cell(row=current_row, column=8, value=int(game["homerun"]))
                ws.cell(row=current_row, column=9, value=int(game["walk_count"]))
                ws.cell(row=current_row, column=10, value=int(game["strike_out"]))
                ws.cell(row=current_row, column=11, value=int(game["stolen_base"]))

                cum_avg = game.get("cum_avg")
                cum_ops = game.get("cum_ops")
                ws.cell(row=current_row, column=12, value=cum_avg)
                ws.cell(row=current_row, column=13, value=cum_ops)
                if cum_avg is not None:
                    ws.cell(row=current_row, column=12).number_format = "0.000"
                if cum_ops is not None:
                    ws.cell(row=current_row, column=13).number_format = "0.000"

                current_row += 1

            data_end = current_row - 1
            if data_end >= data_start:
                self._apply_borders(ws, data_start - 1, data_end, 1, NUM_COLS)

            chart_series_info.append({
                "player_name": player_name,
                "data_start": data_start,
                "data_end": data_end,
            })

            current_row += 1  # blank row between players

        # Column widths
        ws.column_dimensions["A"].width = 13
        ws.column_dimensions["B"].width = 16
        ws.column_dimensions["C"].width = 5
        for col in range(4, 12):
            ws.column_dimensions[get_column_letter(col)].width = 6
        ws.column_dimensions[get_column_letter(12)].width = 9
        ws.column_dimensions[get_column_letter(13)].width = 9

        # Line chart: 累積打率推移（上位5選手 = 出場試合数順）
        top5 = sorted(
            chart_series_info,
            key=lambda x: x["data_end"] - x["data_start"],
            reverse=True,
        )[:5]

        if top5:
            chart = LineChart()
            chart.title = "累積打率推移"
            chart.y_axis.title = "打率"
            chart.x_axis.title = "試合数"
            chart.width = 22
            chart.height = 14

            for info in top5:
                if info["data_end"] >= info["data_start"]:
                    data_ref = Reference(
                        ws,
                        min_col=12,
                        min_row=info["data_start"],
                        max_row=info["data_end"],
                    )
                    chart.add_data(data_ref, titles_from_data=False)
                    chart.series[-1].title = SeriesLabel(v=info["player_name"])

            ws.add_chart(chart, f"O2")

        self.logger.info("Sheet 3 書き込み完了: %d 選手", len(player_rows))

    # ------------------------------------------------------------------ save

    def save(self) -> None:
        """Excelファイルを保存"""
        self.wb.save(self.output_path)
        self.logger.info("✓ Excelファイル保存: %s", self.output_path)
