"""試合スコア・打撃成績をExcelに書き込むリポジトリ"""
import logging
import re
from copy import copy
from openpyxl import load_workbook
from model.game_record import GameRecord
from utils.game_input_parser import format_date_ja

logger = logging.getLogger("app.score_writer")

# 回数ヘッダー行の列位置（C列=3 が1回目）
INNING_START_COL = 3   # C列
MIN_START_ROW = 487    # この行未満には書き込まない（486行目以前は既存データ）
INNING_END_COL = 11    # K列（9回まで）
EXTRA_INNING_COL = 12  # L列（10回・延長戦用。9回終了の場合は空欄）
TOTAL_COL = 13         # M列（L列を挟んだ1列空きの次）
PITCHER_COL = 3        # C列（N+4行目にP➡ピッチャー情報）
BATTING_STATS_COL = 8  # H列（打撃成績）
DATE_COL = 1           # A列


class ScoreExcelWriter:
    """試合スコアを Excel に書き込むクラス"""

    def is_date_already_written(self, ws, date_str: str, game_number: int = 1) -> bool:
        """
        Excelシートをスキャンし、指定日付の game_number 番目の試合がすでに書き込まれているか確認する。

        仕様変更後、2試合目以降は A 列に日付を書かないため、回数ヘッダー行（C=1, D=2）の
        出現数で書き込み済み試合数を判定する。

        Args:
            ws: openpyxl ワークシート
            date_str: "YYYY-MM-DD" 形式の日付文字列
            game_number: 試合番号（1始まり）

        Returns:
            True: game_number 回分すでに書き込まれている / False: まだ書き込まれていない
        """
        date_ja = format_date_ja(date_str)

        # A 列から日付行を探す
        date_rows = []
        for row in ws.iter_rows(min_col=DATE_COL, max_col=DATE_COL, values_only=False):
            cell = row[0]
            if isinstance(cell.value, str) and cell.value == date_ja:
                date_rows.append(cell.row)

        if not date_rows:
            return False

        if game_number == 1:
            return True  # 日付行が1件以上あれば1試合目は書き込み済み

        # 2試合目以降: 日付行の先頭から回数ヘッダー行（C=1, D=2）の数を数える
        first_date_row = date_rows[0]
        block_count = 0
        search_end = first_date_row + game_number * 8
        for r in range(first_date_row, search_end + 1):
            c_val = ws.cell(row=r, column=3).value
            d_val = ws.cell(row=r, column=4).value
            if c_val == 1 and d_val == 2:
                block_count += 1

        return block_count >= game_number

    def find_next_row(self, ws) -> int:
        """
        シートの最後のデータ行を検出し、その +2 行（空行を挟む）を返す。
        A列またはC列に値がある最終行を探す。
        """
        last_data_row = 1
        for row in ws.iter_rows(min_row=1, values_only=False):
            for cell in row:
                if cell.value is not None and cell.column in (1, 3):
                    if cell.row > last_data_row:
                        last_data_row = cell.row
        # 既存データの最終行 + 2（空行1行を挟む）。MIN_START_ROW を下限とする
        return max(last_data_row + 2, MIN_START_ROW)

    def write_game(
        self,
        excel_path: str,
        sheet_name: str,
        game: GameRecord,
        start_row: int = None,
        dry_run: bool = False,
    ) -> int | None:
        """
        試合ブロック（6行）をExcelに書き込む。

        Args:
            excel_path: 書き込み先Excelファイルパス
            sheet_name: シート名
            game: GameRecord オブジェクト
            start_row: 書き込み開始行（None の場合は自動検出）
            dry_run: True の場合は書き込まず内容のみ表示

        Returns:
            書き込んだ開始行番号。同じ日付がすでに存在する場合は None を返す。
        """
        # .xlsm のみ keep_vba=True（.xlsx に適用するとファイルが壊れる）
        keep_vba = excel_path.lower().endswith(".xlsm")
        wb = load_workbook(excel_path, keep_vba=keep_vba)

        if sheet_name not in wb.sheetnames:
            raise ValueError(f"シート '{sheet_name}' が見つかりません。利用可能: {wb.sheetnames}")

        ws = wb[sheet_name]

        # 既存日付チェック（dry_run 時もスキップ対象とする）
        if self.is_date_already_written(ws, game.date, game.game_number):
            logger.info(
                "スキップ: %s 第%d試合はすでにExcelに書き込まれています",
                game.date, game.game_number,
            )
            return None

        if start_row is None:
            start_row = self.find_next_row(ws)

        logger.info("書き込み開始行: %d", start_row)

        # ---- 書き込み内容を構築 ----
        date_ja = format_date_ja(game.date)
        batting_stats_str = _format_batting_stats(game)

        rows_data = _build_rows(game, date_ja, batting_stats_str, start_row)

        if dry_run:
            _print_dry_run(rows_data, start_row, game, date_ja, batting_stats_str)
            return start_row

        # ---- 既存ブロックの書式をコピー ----
        refs = _find_style_refs(ws, start_row)
        if refs:
            row_map = [
                (start_row + 0, refs.get("date")),           # 日付行
                (start_row + 1, refs.get("inning_header")),  # 回数ヘッダー
                (start_row + 2, refs.get("team_score")),     # 自チームスコア
                (start_row + 3, refs.get("team_score")),     # 相手チームスコア
                (start_row + 4, refs.get("pitcher")),        # ピッチャー/打撃成績
            ]
            for target_row, source_row in row_map:
                if source_row is not None:
                    _copy_row_styles(ws, source_row, target_row)

        # ---- 仕様3: 打撃成績のフォントを投手成績と同じにする ----
        pitcher_cell = ws.cell(row=start_row + 4, column=PITCHER_COL)
        batting_cell = ws.cell(row=start_row + 4, column=BATTING_STATS_COL)
        if pitcher_cell.has_style:
            batting_cell.font = copy(pitcher_cell.font)

        # ---- 実際に書き込む ----
        for row_offset, col, value in rows_data:
            ws.cell(row=start_row + row_offset, column=col, value=value)

        # ---- 仕様1: チーム名行の A・B 列を結合 ----
        ws.merge_cells(
            start_row=start_row + 2, start_column=1,
            end_row=start_row + 2, end_column=2,
        )
        ws.merge_cells(
            start_row=start_row + 3, start_column=1,
            end_row=start_row + 3, end_column=2,
        )

        wb.save(excel_path)
        logger.info("Excelに書き込み完了: %s (行 %d〜%d)", excel_path, start_row, start_row + 5)

        return start_row


def _build_rows(game: GameRecord, date_ja: str, batting_stats_str: str, start_row: int) -> list:
    """
    書き込むセルデータを (row_offset, col, value) のリストで返す。
    row_offset は start_row からの相対行。
    """
    entries = []

    # N行目: 日付（A列）+ グランド名（E列）
    # 仕様4: 1日2試合目以降は日付・グランド名を記載しない
    if game.game_number == 1:
        entries.append((0, DATE_COL, date_ja))
        if game.venue:
            entries.append((0, 5, game.venue))  # E列 = グランド名（仕様2）

    # N+1行目: 回数ヘッダー（C〜K列）+ 「計」（M列）
    for i in range(9):
        entries.append((1, INNING_START_COL + i, i + 1))
    entries.append((1, TOTAL_COL, "計"))

    # N+2行目: 自チームのスコア（A列=チーム名, C〜K列=各回, L列=空き（延長戦用）, M列=合計）
    # 仕様5: 表の攻撃行はXを記載しない
    our_row = start_row + 2
    entries.append((2, DATE_COL, game.our_team))
    for i, score in enumerate(_pad_innings(game.our_innings)):
        entries.append((2, INNING_START_COL + i, score))
    our_total_val = _resolve_total(game.our_total, game.our_innings)
    if isinstance(our_total_val, str):
        entries.append((2, TOTAL_COL, our_total_val))
    else:
        entries.append((2, TOTAL_COL, f"=SUM(C{our_row}:L{our_row})"))

    # N+3行目: 相手チームのスコア（裏の攻撃行）
    # 仕様5: サヨナラ勝ち以外はXを記載しない
    opp_row = start_row + 3
    opp_name = game.opp_team if game.opp_team else game.opponent
    entries.append((3, DATE_COL, opp_name))
    for i, score in enumerate(_pad_innings(game.opp_innings)):
        entries.append((3, INNING_START_COL + i, score))
    opp_total_val = _resolve_total(game.opp_total, game.opp_innings)
    if isinstance(opp_total_val, str):
        entries.append((3, TOTAL_COL, opp_total_val))
    else:
        entries.append((3, TOTAL_COL, f"=SUM(C{opp_row}:L{opp_row})"))

    # N+4行目: ピッチャー情報（C列: "P➡{pitcher}"）、打撃成績（H列）
    if game.pitcher:
        entries.append((4, PITCHER_COL, f"P➡{game.pitcher}"))
    entries.append((4, BATTING_STATS_COL, batting_stats_str))

    # N+5行目: 空行（何も書かない）

    return entries


def _pad_innings(innings: list, length: int = 9) -> list:
    """9要素になるよう None で埋める。innings が None の場合は全て None を返す"""
    if innings is None:
        return [None] * length
    padded = list(innings) + [None] * length
    return padded[:length]


def _last_inning_is_number_x(innings) -> bool:
    """
    最後の非空イニングが '1X' のような 数値+X 形式かどうか判定する。
    例: innings[-1] = "1X" → True
         innings[-1] = "X"  → False
    """
    if not innings:
        return False
    for val in reversed(innings):
        if val is None:
            continue
        if isinstance(val, str) and re.match(r'^\d+[Xx]', val):
            return True
        break
    return False


def _resolve_total(total, innings):
    """
    合計値とイニングリストから Excel に書き込む値を決定する。

    - 最後の非空イニングが '1X' のような 数値+X → total を文字列のまま返す（例: "2X"）
      → イニングに得点を含むサヨナラのため、計欄にも "2X" を表示する
    - それ以外で total が "5X" → X を除いた数値(int)を返す → SUM 数式を使用
      → イニングが純粋な "X"（得点なし）のため、計欄は数値のみ表示する
    - total が数値 → そのまま返す
    """
    if not isinstance(total, str):
        return total
    if _last_inning_is_number_x(innings):
        return total  # "2X" のまま返す
    stripped = total.rstrip("Xx")
    if stripped.lstrip("-").isdigit():
        return int(stripped)
    return total


def _format_batting_stats(game: GameRecord) -> str:
    """
    打撃成績セル用文字列を生成する。
    例: "2塁打：選手A、選手B、選手C　　3塁打：選手D"
    種別間は全角スペース2つ。該当者がいない種別は省略。
    全員いない場合は "２塁打以上なし"。
    """
    parts = []

    if game.doubles:
        parts.append(f"2塁打：{'、'.join(game.doubles)}")
    if game.triples:
        parts.append(f"3塁打：{'、'.join(game.triples)}")
    if game.homeruns:
        parts.append(f"本塁打：{'、'.join(game.homeruns)}")

    if not parts:
        return "２塁打以上なし"

    return "\u3000\u3000".join(parts)  # 全角スペース2つで結合


def _find_style_refs(ws, start_row: int) -> dict:
    """
    start_row の直前ブロックを逆検索し、各行タイプのテンプレート行番号を返す。

    検索ロジック:
    - 回数ヘッダー行: C列=1 かつ D列=2 の行
    - チームスコア行: 回数ヘッダーの1行下
    - ピッチャー行: 回数ヘッダーの3行下
    - 日付行: A列に '年' を含む文字列の行

    Returns:
        {'date': int|None, 'inning_header': int|None,
         'team_score': int|None, 'pitcher': int|None}
    """
    refs = {"date": None, "inning_header": None, "team_score": None, "pitcher": None}
    search_limit = max(1, start_row - 50)

    for row_num in range(start_row - 2, search_limit - 1, -1):
        a_val = ws.cell(row=row_num, column=1).value
        c_val = ws.cell(row=row_num, column=3).value
        d_val = ws.cell(row=row_num, column=4).value

        # 回数ヘッダー行: C=1, D=2
        if refs["inning_header"] is None and c_val == 1 and d_val == 2:
            refs["inning_header"] = row_num
            if refs["team_score"] is None:
                refs["team_score"] = row_num + 1
            if refs["pitcher"] is None:
                refs["pitcher"] = row_num + 3

        # 日付行: A列に '年' を含む文字列
        if refs["date"] is None and isinstance(a_val, str) and "年" in a_val:
            refs["date"] = row_num

        if all(v is not None for v in refs.values()):
            break

    return refs


def _copy_row_styles(ws, source_row: int, target_row: int) -> None:
    """
    source_row の全セルのスタイル（フォント・罫線・塗りつぶし・配置・表示形式）と
    行の高さを target_row にコピーする。
    """
    # セルスタイル
    for col in range(1, ws.max_column + 1):
        src = ws.cell(row=source_row, column=col)
        tgt = ws.cell(row=target_row, column=col)
        if src.has_style:
            tgt.font = copy(src.font)
            tgt.border = copy(src.border)
            tgt.fill = copy(src.fill)
            tgt.number_format = src.number_format
            tgt.protection = copy(src.protection)
            tgt.alignment = copy(src.alignment)

    # 行の高さ
    src_height = ws.row_dimensions[source_row].height
    if src_height is not None:
        ws.row_dimensions[target_row].height = src_height


def _print_dry_run(rows_data, start_row, game, date_ja, batting_stats_str):
    """--dry-run 時の確認表示"""
    print("=" * 60)
    print(f"[DRY RUN] 書き込み内容プレビュー (開始行: {start_row})")
    print("=" * 60)

    our_innings_str = ", ".join("" if s is None else str(s) for s in _pad_innings(game.our_innings))
    opp_innings_str = ", ".join("" if s is None else str(s) for s in _pad_innings(game.opp_innings))
    opp_name = game.opp_team if game.opp_team else game.opponent

    if game.game_number == 1:
        venue_str = f"  E列 = {game.venue}" if game.venue else ""
        print(f"  行 {start_row}    (N)  : A列 = {date_ja}{venue_str}")
    else:
        print(f"  行 {start_row}    (N)  : (日付・グランド名なし: 第{game.game_number}試合目)")

    print(f"  行 {start_row+1}  (N+1): 回数ヘッダー (1〜9, 計)")

    our_total_val = _resolve_total(game.our_total, game.our_innings)
    our_total_str = our_total_val if isinstance(our_total_val, str) else f"=SUM(C{start_row+2}:L{start_row+2})"
    print(f"  行 {start_row+2}  (N+2): [A:B結合] {game.our_team} [{our_innings_str}] 計:{our_total_str}")

    opp_total_val = _resolve_total(game.opp_total, game.opp_innings)
    opp_total_str = opp_total_val if isinstance(opp_total_val, str) else f"=SUM(C{start_row+3}:L{start_row+3})"
    sayonara_mark = " [サヨナラ]" if game.is_sayonara else ""
    print(f"  行 {start_row+3}  (N+3): [A:B結合] {opp_name} [{opp_innings_str}] 計:{opp_total_str}{sayonara_mark}")

    print(f"  行 {start_row+4}  (N+4): P➡{game.pitcher}")
    print(f"           打撃成績(投手と同フォント): {batting_stats_str}")
    print(f"  行 {start_row+5}  (N+5): (空行)")
    print("=" * 60)
