"""ScoreExcelWriter（スコアの Excel 書込）のテスト"""
import os
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from io import StringIO

from openpyxl import Workbook, load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from model.game_record import GameRecord
from repository import score_excel_writer as sew
from repository.score_excel_writer import ScoreExcelWriter
from utils.game_input_parser import format_date_ja

SHEET = "スコア"
DATE = "2026-07-21"  # 火曜日


def _game(**kwargs) -> GameRecord:
    base = dict(
        date=DATE, opponent="相手校", game_number=1, venue="第一グラウンド",
        pitcher="投手A（6回）、投手B（3回）",
        our_innings=[2, 0, 4, 1, 1, 0, 2, 1, 2], opp_innings=[3, 0, 2, 0, 0, 2, 0, 0, 1],
        our_total=13, opp_total=8,
        our_team="自チーム", opp_team="相手校",
        doubles=["選手A", "選手B"], triples=["選手D"], homeruns=[],
    )
    base.update(kwargs)
    return GameRecord(**base)


class TestHelpers(unittest.TestCase):

    def test_pad_innings(self):
        self.assertEqual(sew._pad_innings([1, 2]), [1, 2] + [None] * 7)
        self.assertEqual(sew._pad_innings(None), [None] * 9)
        self.assertEqual(len(sew._pad_innings(list(range(12)))), 9)

    def test_resolve_total_numeric(self):
        self.assertEqual(sew._resolve_total(5, [1, 4]), 5)

    def test_resolve_total_x_without_score_in_last_inning_is_numeric(self):
        # 最終回が純粋な "X"（得点なし）: 計は数値だけ
        self.assertEqual(sew._resolve_total("5X", [1, 4, "X"]), 5)

    def test_resolve_total_x_with_score_in_last_inning_stays_string(self):
        # 最終回が "1X": 計も "5X" のまま
        self.assertEqual(sew._resolve_total("5X", [1, 3, "1X"]), "5X")

    def test_format_batting_stats(self):
        text = sew._format_batting_stats(_game())
        self.assertEqual(text, "2塁打：選手A、選手B　　3塁打：選手D")

    def test_format_batting_stats_none(self):
        text = sew._format_batting_stats(_game(doubles=[], triples=[], homeruns=[]))
        self.assertEqual(text, "２塁打以上なし")


class TestBuildRows(unittest.TestCase):

    def _cells(self, game, start_row=490):
        rows = sew._build_rows(game, format_date_ja(game.date), "stats", start_row)
        return {(r, c): v for r, c, v in rows}

    def test_first_game_has_date_and_venue(self):
        cells = self._cells(_game())
        self.assertEqual(cells[(0, sew.DATE_COL)], "2026年7月21日（火）")
        self.assertEqual(cells[(0, 5)], "第一グラウンド")

    def test_second_game_has_no_date(self):
        cells = self._cells(_game(game_number=2))
        self.assertNotIn((0, sew.DATE_COL), cells)

    def test_team_names_and_scores(self):
        cells = self._cells(_game())
        self.assertEqual(cells[(2, sew.DATE_COL)], "自チーム")
        self.assertEqual(cells[(3, sew.DATE_COL)], "相手校")
        self.assertEqual(cells[(2, sew.INNING_START_COL)], 2)
        self.assertEqual(cells[(3, sew.INNING_START_COL + 2)], 2)

    def test_total_is_sum_formula_for_numeric_total(self):
        cells = self._cells(_game(), start_row=490)
        self.assertEqual(cells[(2, sew.TOTAL_COL)], "=SUM(C492:L492)")
        self.assertEqual(cells[(3, sew.TOTAL_COL)], "=SUM(C493:L493)")

    def test_pitcher_and_batting_stats(self):
        cells = self._cells(_game())
        self.assertEqual(cells[(4, sew.PITCHER_COL)], "P➡投手A（6回）、投手B（3回）")
        self.assertEqual(cells[(4, sew.BATTING_STATS_COL)], "stats")

    def test_blank_score_when_innings_missing(self):
        cells = self._cells(_game(our_innings=None, opp_innings=None, our_total=None, opp_total=None))
        self.assertIsNone(cells[(2, sew.INNING_START_COL)])


class TestWriteGame(unittest.TestCase):

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path = os.path.join(self.tmp.name, "score.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = SHEET
        ws.cell(row=1, column=1, value="既存データ")
        wb.save(self.path)
        self.writer = ScoreExcelWriter()

    def _sheet(self):
        return load_workbook(self.path)[SHEET]

    def test_first_write_starts_at_row_487(self):
        # 486行目以前は既存データのため、487行目より前には書き込まない
        start = self.writer.write_game(self.path, SHEET, _game())
        self.assertEqual(start, 487)

        ws = self._sheet()
        self.assertEqual(ws.cell(row=start, column=1).value, "2026年7月21日（火）")
        self.assertEqual(ws.cell(row=start + 2, column=1).value, "自チーム")
        self.assertEqual(ws.cell(row=start + 3, column=1).value, "相手校")
        self.assertEqual(ws.cell(row=start + 2, column=sew.TOTAL_COL).value, f"=SUM(C{start + 2}:L{start + 2})")

    def test_existing_data_is_kept(self):
        self.writer.write_game(self.path, SHEET, _game())
        self.assertEqual(self._sheet().cell(row=1, column=1).value, "既存データ")

    def test_team_name_cells_are_merged(self):
        start = self.writer.write_game(self.path, SHEET, _game())
        merged = {str(r) for r in self._sheet().merged_cells.ranges}
        self.assertIn(f"A{start + 2}:B{start + 2}", merged)
        self.assertIn(f"A{start + 3}:B{start + 3}", merged)

    def test_same_game_is_skipped(self):
        self.writer.write_game(self.path, SHEET, _game())
        self.assertIsNone(self.writer.write_game(self.path, SHEET, _game()))

    def test_second_game_of_the_day_is_written_after_first(self):
        first = self.writer.write_game(self.path, SHEET, _game())
        second = self.writer.write_game(self.path, SHEET, _game(game_number=2))
        self.assertIsNotNone(second)
        self.assertGreater(second, first)
        # 2試合目は日付を書かない
        self.assertIsNone(self._sheet().cell(row=second, column=1).value)

    def test_dry_run_does_not_modify_file(self):
        before = os.path.getmtime(self.path)
        with redirect_stdout(StringIO()):
            start = self.writer.write_game(self.path, SHEET, _game(), dry_run=True)
        self.assertEqual(start, 487)
        self.assertEqual(os.path.getmtime(self.path), before)
        self.assertIsNone(self._sheet().cell(row=start, column=1).value)

    def test_unknown_sheet_raises(self):
        with self.assertRaises(ValueError):
            self.writer.write_game(self.path, "存在しないシート", _game())

    def test_explicit_start_row(self):
        start = self.writer.write_game(self.path, SHEET, _game(), start_row=600)
        self.assertEqual(start, 600)
        self.assertEqual(self._sheet().cell(row=600, column=1).value, "2026年7月21日（火）")


if __name__ == "__main__":
    unittest.main()
